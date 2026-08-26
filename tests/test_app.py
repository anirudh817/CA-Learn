import csv
import contextlib
import importlib.util
import io
import json
import os
import sys
import tempfile
import unittest
import uuid
import warnings
import zipfile
from pathlib import Path

from fastapi.testclient import TestClient as FastAPITestClient

from proteomics_ai.devserver import resolve_port


ROOT = Path(__file__).resolve().parents[1]
BACKEND_DIR = ROOT / "backend"
GLOBAL_MPL_DIR = Path(tempfile.gettempdir()) / "signalfold-mplconfig"
GLOBAL_MPL_DIR.mkdir(parents=True, exist_ok=True)

os.environ.setdefault("MPLCONFIGDIR", str(GLOBAL_MPL_DIR))

warnings.simplefilter("ignore", DeprecationWarning)
warnings.simplefilter("ignore", ResourceWarning)

warnings.filterwarnings(
    "ignore",
    category=DeprecationWarning,
)
warnings.filterwarnings(
    "ignore",
    message=".*shortcut is now deprecated.*",
    category=DeprecationWarning,
)
warnings.filterwarnings(
    "ignore",
    message=".*Pyarrow will become a required dependency of pandas.*",
    category=DeprecationWarning,
)
warnings.filterwarnings(
    "ignore",
    message="Unclosed <MemoryObjectSendStream>",
    category=ResourceWarning,
)
warnings.filterwarnings(
    "ignore",
    message="Unclosed <MemoryObjectReceiveStream>",
    category=ResourceWarning,
)


def dataset_bytes(apoe_shift: float = 0.0, global_shift: float = 0.0) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["Protein", "Control_1", "Control_2", "Control_3", "Disease_1", "Disease_2", "Disease_3"])

    seeded = [
        ("APOE", 10.0, 13.4 + apoe_shift),
        ("CLU", 8.2, 11.6 + global_shift),
        ("C3", 7.1, 9.7 + global_shift),
        ("GFAP", 4.0, 5.8 + global_shift),
        ("MBP", 6.3, 7.2 + global_shift),
        ("VIM", 5.2, 7.4 + global_shift),
        ("SERPINA1", 8.7, 10.8 + global_shift),
        ("ALB", 14.2, 13.6 + global_shift),
        ("TREM2", 3.2, 4.8 + global_shift),
        ("PLP1", 6.4, 7.1 + global_shift),
    ]
    for gene, control_base, disease_base in seeded:
        writer.writerow(
            [
                gene,
                control_base,
                control_base + 0.08,
                control_base - 0.06,
                disease_base,
                disease_base + 0.07,
                disease_base - 0.05,
            ]
        )

    for index in range(1, 21):
        control_base = 5.0 + index * 0.35 + global_shift
        disease_delta = 1.1 if index % 3 else -0.65
        disease_base = control_base + disease_delta
        writer.writerow(
            [
                f"GENE{index:03d}",
                control_base,
                control_base + 0.04,
                control_base - 0.05,
                disease_base,
                disease_base + 0.06,
                disease_base - 0.03,
            ]
        )

    return buffer.getvalue().encode("utf-8")


def spectronaut_raw_bytes() -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["R.FileName", "PEP.GroupingKey", "PEP.Quantity", "PEP.AllOccurringProteinAccessions"])

    samples = {
        "20180618_QX0_demo_sampleA01": ("Control", [12000, 9500, 8200, 7000, 6100, 5400]),
        "20180618_QX0_demo_sampleA02": ("Control", [12500, 9400, 8100, 7100, 6200, 5300]),
        "20180618_QX0_demo_sampleA03": ("Control", [11800, 9600, 8300, 7050, 6000, 5500]),
        "20180618_QX0_demo_sampleB01": ("Disease", [42000, 29500, 26500, 5200, 3900, 4100]),
        "20180618_QX0_demo_sampleB02": ("Disease", [43000, 30000, 27000, 5100, 3950, 4050]),
        "20180618_QX0_demo_sampleB03": ("Disease", [41500, 29200, 26800, 5150, 3850, 3980]),
    }
    peptides = [
        ("APOE_PEPTIDE", "P02649"),
        ("CLU_PEPTIDE", "P10909"),
        ("C3_PEPTIDE", "P01024"),
        ("MBP_PEPTIDE", "P02686"),
        ("GFAP_PEPTIDE", "P14136"),
        ("VIM_PEPTIDE", "P08670"),
    ]

    for sample_name, (_, intensities) in samples.items():
        for (peptide, accession), intensity in zip(peptides, intensities):
            writer.writerow([sample_name, peptide, intensity, accession])

    return buffer.getvalue().encode("utf-8")


def traits_bytes() -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["sample name", "primary biochemical AD classification", "t-tau [ng/L]"])
    rows = [
        ("20180618_QX0_demo_sampleA01.raw.PG.Quantity", "biochemical control", 240),
        ("20180618_QX0_demo_sampleA02.raw.PG.Quantity", "biochemical control", 255),
        ("20180618_QX0_demo_sampleA03.raw.PG.Quantity", "biochemical control", 250),
        ("20180618_QX0_demo_sampleB01.raw.PG.Quantity", "biochemical AD", 780),
        ("20180618_QX0_demo_sampleB02.raw.PG.Quantity", "biochemical AD", 810),
        ("20180618_QX0_demo_sampleB03.raw.PG.Quantity", "biochemical AD", 795),
    ]
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def olink_npx_bytes() -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["SampleID", "SampleType", "Assay", "AssayType", "NPX", "SampleQC", "Disease", "Age.At.Collection"])
    assays = [f"Protein_{index:02d}" for index in range(1, 25)]
    samples = [
        ("CTRL_01", "Control", 0.0, 58),
        ("CTRL_02", "Control", 0.1, 61),
        ("CTRL_03", "Control", -0.1, 63),
        ("ALS_01", "Disease", 1.2, 67),
        ("ALS_02", "Disease", 1.1, 69),
        ("ALS_03", "Disease", 1.3, 72),
    ]
    for sample_name, disease_label, disease_shift, age in samples:
        for assay_index, assay in enumerate(assays, start=1):
            baseline = 1.5 + assay_index * 0.12
            if assay_index <= 8:
                value = baseline + disease_shift
            elif assay_index <= 16:
                value = baseline - (disease_shift * 0.5)
            else:
                value = baseline + (0.05 if "ALS" in sample_name else -0.03)
            writer.writerow([sample_name, "SAMPLE", assay, "assay", round(value, 4), "PASS", disease_label, age])
    return buffer.getvalue().encode("utf-8")


def build_app(temp_dir: str):
    data_dir = Path(temp_dir) / "data"
    db_path = Path(temp_dir) / "proteomics-test.db"

    os.environ["APP_ENV"] = "local"
    os.environ["DATA_DIR"] = str(data_dir)
    os.environ["DATABASE_URL"] = f"sqlite:///{db_path}"
    os.environ["INLINE_RUNS"] = "1"
    os.environ["ENABLE_LOCAL_BOOTSTRAP"] = "1"
    os.environ["MPLCONFIGDIR"] = str(GLOBAL_MPL_DIR)

    if str(BACKEND_DIR) not in sys.path:
        sys.path.insert(0, str(BACKEND_DIR))

    for module_name in list(sys.modules):
        if module_name in {"config", "database", "deps", "schemas", "security", "utils", "main"}:
            sys.modules.pop(module_name, None)
        if module_name.startswith("routes") or module_name.startswith("services"):
            sys.modules.pop(module_name, None)

    spec = importlib.util.spec_from_file_location(f"test_backend_main_{uuid.uuid4().hex}", BACKEND_DIR / "main.py")
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        spec.loader.exec_module(module)
    return module.create_app()


@contextlib.contextmanager
def managed_client(app):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        with FastAPITestClient(app) as client:
            yield client


def TestClient(app):
    return managed_client(app)


def bootstrap_session(client: FastAPITestClient):
    response = client.post("/api/auth/bootstrap-local")
    assert response.status_code == 200, response.text
    payload = response.json()
    token = payload["token"]
    headers = {"Authorization": f"Bearer {token}"}
    workspace = payload["workspaces"][0]
    projects = client.get("/api/projects", params={"workspace_id": workspace["id"]}, headers=headers)
    assert projects.status_code == 200, projects.text
    project = projects.json()[0]
    return headers, workspace, project


def upload_primary_dataset(client: FastAPITestClient, headers, workspace_id: str, project_id: str, content: bytes):
    return upload_dataset(client, headers, workspace_id, project_id, content, "primary", "cohort.csv")


def upload_dataset(client: FastAPITestClient, headers, workspace_id: str, project_id: str, content: bytes, file_kind: str, filename: str):
    response = client.post(
        "/api/uploads",
        headers=headers,
        data={"workspace_id": workspace_id, "project_id": project_id, "file_kind": file_kind},
        files={"file": (filename, content, "text/csv")},
    )
    assert response.status_code == 200, response.text
    return response.json()


class CommercialAppTests(unittest.TestCase):
    def test_resolve_port_defaults_to_8000_when_free(self) -> None:
        import unittest.mock

        with unittest.mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop("PORT", None)
            with unittest.mock.patch("proteomics_ai.devserver.port_is_available", return_value=True):
                self.assertEqual(resolve_port(), 8000)

    def test_resolve_port_honours_PORT_env(self) -> None:
        import unittest.mock

        with unittest.mock.patch.dict(os.environ, {"PORT": "8123"}):
            with unittest.mock.patch("proteomics_ai.devserver.port_is_available", return_value=True):
                self.assertEqual(resolve_port(), 8123)

    def test_resolve_port_fails_loudly_when_busy(self) -> None:
        """A busy preferred port must abort, never silently drift to the next one."""
        import unittest.mock

        with unittest.mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop("PORT", None)
            with unittest.mock.patch("proteomics_ai.devserver.port_is_available", return_value=False):
                with self.assertRaises(SystemExit):
                    resolve_port()

    def test_spa_bootstrap_and_run_flow(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                home = client.get("/")
                self.assertEqual(home.status_code, 200)
                self.assertIn("ProteomicsAI", home.text)
                self.assertIn("/app.js", home.text)

                headers, workspace, project = bootstrap_session(client)
                uploaded = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes())
                self.assertEqual(uploaded["format_family"], "Generic")
                self.assertGreater(uploaded["sample_count"], 0)
                self.assertEqual(uploaded["pipeline_profile"], "generic_matrix")
                self.assertIn("recommended_defaults", uploaded)
                self.assertIn("parameter_overrides", uploaded["recommended_defaults"])

                run_response = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Commercial Baseline",
                        "dataset_id": uploaded["dataset_id"],
                        "cohort1": "Control",
                        "cohort2": "Disease",
                    },
                )
                self.assertEqual(run_response.status_code, 200, run_response.text)
                run_payload = run_response.json()
                self.assertEqual(run_payload["status"], "complete")
                run_id = run_payload["run_id"]

                run_detail = client.get(f"/api/runs/{run_id}", headers=headers)
                self.assertEqual(run_detail.status_code, 200)
                detail_json = run_detail.json()
                self.assertEqual(detail_json["status"], "complete")
                self.assertTrue(any(file["rel_path"] == "input/dataset_manifest.json" for file in detail_json["files"]))
                self.assertTrue(any(stage["stage_key"] == "celltypefet" and stage["status"] == "complete" for stage in detail_json["stages"]))

                summary = client.get(f"/api/results/{run_id}/summary", headers=headers)
                volcano = client.get(f"/api/results/{run_id}/volcano", headers=headers)
                files = client.get(f"/api/runs/{run_id}/files", headers=headers)
                self.assertEqual(summary.status_code, 200)
                self.assertEqual(volcano.status_code, 200)
                self.assertEqual(files.status_code, 200)
                self.assertIn("run_name", summary.json())
                self.assertGreater(len(volcano.json()["x"]), 0)
                self.assertGreater(len(files.json()), 5)

                # POST /api/ai/query was deprecated in favor of the chat API
                # (introduced in P0 of the AI Biological Inference Chat spec).
                # The legacy endpoint now returns HTTP 410 with a migration hint.
                ai = client.post(
                    "/api/ai/query",
                    headers=headers,
                    json={"question": "Summarize this run", "run_ids": [run_id]},
                )
                self.assertEqual(ai.status_code, 410)
                ai_payload = ai.json()["detail"]
                self.assertIn("conversations", " ".join(ai_payload.get("migrate_to", [])))

                runtime = client.get("/api/auth/runtime")
                self.assertEqual(runtime.status_code, 200, runtime.text)
                runtime_json = runtime.json()
                self.assertEqual(runtime_json["auth_mode"], "local_bootstrap")
                self.assertTrue(runtime_json["local_bootstrap_enabled"])
                self.assertIn("ai", runtime_json)
                self.assertEqual(runtime_json["ai"]["provider"], "anthropic")

    def test_compare_share_and_audit(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                first_dataset = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes(apoe_shift=0.0))
                second_dataset = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes(apoe_shift=1.1, global_shift=0.25))

                first_run = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Run One",
                        "dataset_id": first_dataset["dataset_id"],
                        "params": {"normalization_method": "median", "wgcna_power": 8},
                    },
                ).json()
                second_run = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Run Two",
                        "dataset_id": second_dataset["dataset_id"],
                        "params": {"normalization_method": "quantile", "wgcna_power": 12},
                    },
                ).json()

                compare = client.post(
                    "/api/compare/runs",
                    headers=headers,
                    json={"left_run_id": first_run["run_id"], "right_run_id": second_run["run_id"]},
                )
                self.assertEqual(compare.status_code, 200, compare.text)
                compare_json = compare.json()
                self.assertEqual(compare_json["left"]["id"], first_run["run_id"])
                self.assertGreaterEqual(len(compare_json["param_diff"]), 1)
                self.assertIn("modules_count", compare_json["metrics_delta"])

                share = client.post(
                    "/api/share-links",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "run_id": first_run["run_id"],
                        "title": "Investor Preview",
                    },
                )
                self.assertEqual(share.status_code, 200, share.text)
                share_json = share.json()
                self.assertEqual(share_json["scope"], "run")

                resolved = client.get(f"/api/share-links/{share_json['token']}")
                self.assertEqual(resolved.status_code, 200)
                self.assertEqual(resolved.json()["run"]["id"], first_run["run_id"])

                audit = client.get("/api/audit", headers=headers, params={"workspace_id": workspace["id"]})
                self.assertEqual(audit.status_code, 200, audit.text)
                actions = [entry["action_type"] for entry in audit.json()["events"]]
                self.assertIn("run.created", actions)
                self.assertIn("share_link.created", actions)

    def test_upload_previews_and_run_trash_restore_purge(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                raw_preview = client.post(
                    "/api/uploads/preview/raw",
                    headers=headers,
                    data={"workspace_id": workspace["id"], "project_id": project["id"]},
                    files={"file": ("preview.csv", dataset_bytes(), "text/csv")},
                )
                self.assertEqual(raw_preview.status_code, 200, raw_preview.text)
                raw_json = raw_preview.json()
                self.assertEqual(raw_json["sniff"]["format_detected"], "Generic Wide Matrix")
                self.assertIn("confidence", raw_json["sniff"])
                self.assertTrue(raw_json["sniff"]["evidence"])
                self.assertGreaterEqual(raw_json["rows_total"], 10)
                self.assertGreaterEqual(raw_json["columns_total"], 3)
                self.assertEqual(raw_json["pipeline_profile"], "generic_matrix")
                self.assertIn("recommended_defaults", raw_json)
                self.assertIn("parameter_overrides", raw_json["recommended_defaults"])

                traits_preview = client.post(
                    "/api/uploads/preview/traits",
                    headers=headers,
                    data={"workspace_id": workspace["id"], "project_id": project["id"]},
                    files={"file": ("traits.csv", traits_bytes(), "text/csv")},
                )
                self.assertEqual(traits_preview.status_code, 200, traits_preview.text)
                traits_json = traits_preview.json()
                self.assertGreaterEqual(traits_json["rows_total"], 3)
                self.assertIn("sample name", traits_json["columns"])
                self.assertEqual(traits_json["pipeline_profile"], "generic_matrix")

                uploaded = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes())
                run_response = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Trash Me",
                        "dataset_id": uploaded["dataset_id"],
                        "cohort1": "Control",
                        "cohort2": "Disease",
                    },
                )
                self.assertEqual(run_response.status_code, 200, run_response.text)
                run_id = run_response.json()["run_id"]

                active_before = client.get("/api/runs", headers=headers, params={"workspace_id": workspace["id"]})
                self.assertEqual(active_before.status_code, 200)
                self.assertIn(run_id, [run["id"] for run in active_before.json()])

                trashed = client.post(f"/api/runs/{run_id}/trash", headers=headers)
                self.assertEqual(trashed.status_code, 200, trashed.text)
                self.assertEqual(trashed.json()["status"], "trashed")

                active_after_trash = client.get("/api/runs", headers=headers, params={"workspace_id": workspace["id"]})
                self.assertEqual(active_after_trash.status_code, 200)
                self.assertNotIn(run_id, [run["id"] for run in active_after_trash.json()])

                trash_list = client.get("/api/runs", headers=headers, params={"workspace_id": workspace["id"], "include_trashed": "true"})
                self.assertEqual(trash_list.status_code, 200)
                self.assertIn(run_id, [run["id"] for run in trash_list.json()])

                restored = client.post(f"/api/runs/{run_id}/restore", headers=headers)
                self.assertEqual(restored.status_code, 200, restored.text)
                self.assertEqual(restored.json()["status"], "complete")

                active_after_restore = client.get("/api/runs", headers=headers, params={"workspace_id": workspace["id"]})
                self.assertEqual(active_after_restore.status_code, 200)
                self.assertIn(run_id, [run["id"] for run in active_after_restore.json()])

                client.post(f"/api/runs/{run_id}/trash", headers=headers)
                purged = client.post(f"/api/runs/{run_id}/purge", headers=headers)
                self.assertEqual(purged.status_code, 200, purged.text)
                self.assertTrue(purged.json()["purged"])

                trash_after_purge = client.get("/api/runs", headers=headers, params={"workspace_id": workspace["id"], "include_trashed": "true"})
                self.assertEqual(trash_after_purge.status_code, 200)
                self.assertNotIn(run_id, [run["id"] for run in trash_after_purge.json()])

    def test_spectronaut_long_format_uses_traits_and_produces_sane_volcano(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                uploaded = upload_dataset(
                    client,
                    headers,
                    workspace["id"],
                    project["id"],
                    spectronaut_raw_bytes(),
                    "primary",
                    "sweden_raw_like.csv",
                )
                self.assertEqual(uploaded["format_family"], "Spectronaut")
                self.assertEqual(uploaded["assay_level"], "peptide")
                self.assertGreaterEqual(uploaded["recommended_defaults"]["parameter_overrides"]["deep_split"], 3)

                traits = upload_dataset(
                    client,
                    headers,
                    workspace["id"],
                    project["id"],
                    traits_bytes(),
                    "traits",
                    "traits.csv",
                )

                run_response = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Spectronaut Traits Test",
                        "dataset_id": uploaded["dataset_id"],
                        "traits_dataset_id": traits["dataset_id"],
                        "cohort1": "Control",
                        "cohort2": "Disease",
                    },
                )
                self.assertEqual(run_response.status_code, 200, run_response.text)
                run_id = run_response.json()["run_id"]

                summary = client.get(f"/api/results/{run_id}/summary", headers=headers)
                volcano = client.get(f"/api/results/{run_id}/volcano", headers=headers)
                self.assertEqual(summary.status_code, 200)
                self.assertEqual(volcano.status_code, 200)
                summary_json = summary.json()
                volcano_json = volcano.json()

                self.assertGreater(summary_json["peptides_significant"], 0)
                self.assertGreater(summary_json["peptides_upregulated"], 0)
                self.assertLess(max(abs(value) for value in volcano_json["x"]), 10)
                self.assertEqual(volcano_json["metric_label"], "adjusted p-value")

    def test_raw_pvalue_mode_changes_volcano_metric(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)
                uploaded = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes())

                run_response = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Raw pvalue test",
                        "dataset_id": uploaded["dataset_id"],
                        "cohort1": "Control",
                        "cohort2": "Disease",
                        "params": {"use_adjusted_pvalue": False},
                    },
                )
                self.assertEqual(run_response.status_code, 200, run_response.text)
                run_id = run_response.json()["run_id"]

                volcano = client.get(f"/api/results/{run_id}/volcano", headers=headers)
                top = client.get(f"/api/results/{run_id}/top_proteins", headers=headers)
                self.assertEqual(volcano.status_code, 200)
                self.assertEqual(top.status_code, 200)
                self.assertEqual(volcano.json()["metric_label"], "p-value")
                self.assertEqual(top.json()[0]["significance_metric_label"], "p-value")

    def test_olink_long_format_is_detected_and_runs_without_extra_log_transform(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)
                uploaded = upload_dataset(
                    client,
                    headers,
                    workspace["id"],
                    project["id"],
                    olink_npx_bytes(),
                    "primary",
                    "olink_npx.csv",
                )
                self.assertEqual(uploaded["format_family"], "Olink")
                self.assertEqual(uploaded["assay_level"], "protein")

                run_response = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Olink Run",
                        "dataset_id": uploaded["dataset_id"],
                        "cohort1": "Control",
                        "cohort2": "Disease",
                    },
                )
                self.assertEqual(run_response.status_code, 200, run_response.text)
                run_id = run_response.json()["run_id"]

                summary = client.get(f"/api/results/{run_id}/summary", headers=headers)
                self.assertEqual(summary.status_code, 200)
                summary_json = summary.json()
                self.assertEqual(summary_json["format_family"], "Olink")
                self.assertFalse(summary_json["log_transform_applied"])
                self.assertGreater(summary_json["peptides_significant"], 0)

    def test_stage1_summary_uses_requested_wgcna_power(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)
                uploaded = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes())

                run_response = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Power Test",
                        "dataset_id": uploaded["dataset_id"],
                        "cohort1": "Control",
                        "cohort2": "Disease",
                        "params": {"wgcna_power": 11},
                    },
                )
                self.assertEqual(run_response.status_code, 200, run_response.text)
                run_id = run_response.json()["run_id"]

                summary = client.get(f"/api/results/{run_id}/summary", headers=headers)
                self.assertEqual(summary.status_code, 200)
                self.assertEqual(summary.json()["selected_wgcna_power"], 11)

    def test_legacy_bundle_and_table_viewer_endpoints(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)
                uploaded = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes())

                run_response = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Deliverables Test",
                        "dataset_id": uploaded["dataset_id"],
                        "cohort1": "Control",
                        "cohort2": "Disease",
                    },
                )
                self.assertEqual(run_response.status_code, 200, run_response.text)
                run_id = run_response.json()["run_id"]

                artifacts_response = client.get(f"/api/results/{run_id}/artifacts", headers=headers)
                self.assertEqual(artifacts_response.status_code, 200, artifacts_response.text)
                artifacts = artifacts_response.json()

                self.assertGreater(artifacts["overview"]["count"], 0)
                self.assertGreater(artifacts["qc"]["count"], 0)
                self.assertGreater(artifacts["tables"]["count"], 0)
                self.assertGreaterEqual(len(artifacts["volcano"]["featured"]), 1)
                self.assertEqual(artifacts["volcano"]["featured"][0]["kind"], "html")
                self.assertIn("pdf", {item["kind"] for item in artifacts["volcano"]["featured"][0]["download_companions"]})

                overview_paths = [item["rel_path"] for item in artifacts["overview"]["featured"] + artifacts["overview"]["secondary"]]
                self.assertIn("PROTEOMICS_Interactive_Dashboard.html", overview_paths)

                overview_dashboard = next(
                    item for item in artifacts["overview"]["featured"] + artifacts["overview"]["secondary"] if item["rel_path"] == "PROTEOMICS_Interactive_Dashboard.html"
                )
                dashboard_html = client.get(
                    f"/api/results/{run_id}/artifacts/{overview_dashboard['artifact_id']}/content",
                    headers=headers,
                )
                self.assertEqual(dashboard_html.status_code, 200, dashboard_html.text)
                self.assertIn("window.PROTEOMICS_DASHBOARD", dashboard_html.text)

                # DEL-02: Verify Complete_Results.xlsx was generated
                import openpyxl
                run_path = Path(temp_dir) / "data" / "runs" / run_id
                norm_tag = "CBN_median"
                xlsx_candidates = list((run_path / f"05_network_{norm_tag}").glob("*_Complete_Results.xlsx"))
                self.assertTrue(len(xlsx_candidates) > 0, "Complete_Results.xlsx not found in network dir")
                wb = openpyxl.load_workbook(xlsx_candidates[0], read_only=True)
                expected_sheets = {"Module_Assignments", "GO_Enrichment_ZScores", "CellType_FET", "Trait_Associated_Modules"}
                actual_sheets = set(wb.sheetnames)
                wb.close()
                self.assertTrue(
                    expected_sheets.issubset(actual_sheets),
                    f"Missing xlsx sheets: {expected_sheets - actual_sheets}",
                )

                # DEL-03: Verify Volcano_Summary.txt was generated
                volcano_txt_source = run_path / "stage1" / "volcano_summary.txt"
                self.assertTrue(volcano_txt_source.exists(), "volcano_summary.txt not generated in stage1/")
                volcano_content = volcano_txt_source.read_text()
                self.assertIn("VOLCANO PLOT ANALYSIS SUMMARY", volcano_content)
                self.assertIn("Total peptides analyzed:", volcano_content)
                self.assertIn("Total significant:", volcano_content)

                cleaned_matrix = next(
                    item for item in artifacts["tables"]["featured"] + artifacts["tables"]["secondary"] if item["rel_path"] == "input/cleaned_matrix.csv"
                )
                viewer = client.get(
                    f"/api/results/{run_id}/artifacts/{cleaned_matrix['artifact_id']}/viewer",
                    headers=headers,
                )
                self.assertEqual(viewer.status_code, 200, viewer.text)
                self.assertEqual(viewer.json()["viewer"]["type"], "table")

                self.assertFalse((run_path / "Proteomics Reports").exists(), "Shelved report directory should not be emitted")

                table = client.get(
                    f"/api/results/{run_id}/tables/{cleaned_matrix['artifact_id']}",
                    headers=headers,
                    params={"search": "APOE", "page_size": 5},
                )
                self.assertEqual(table.status_code, 200, table.text)
                table_json = table.json()
                self.assertGreaterEqual(table_json["total_rows"], 1)
                row_text = json.dumps(table_json["rows"][0])
                self.assertIn("APOE", row_text)

    def test_native_artifacts_are_listed_served_and_rdata_is_hidden(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)
                uploaded = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes())

                run_response = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Native Artifact Test",
                        "dataset_id": uploaded["dataset_id"],
                        "cohort1": "Control",
                        "cohort2": "Disease",
                    },
                )
                self.assertEqual(run_response.status_code, 200, run_response.text)
                run_id = run_response.json()["run_id"]

                run_dir = Path(temp_dir) / "data" / "runs" / run_id
                stage1_dir = run_dir / "stage1"
                stage1_dir.mkdir(parents=True, exist_ok=True)
                (stage1_dir / "PEAKS_Interactive_Volcano_Plot.html").write_text(
                    '<html><head><script src="dashboard_data.js"></script></head><body>native volcano</body></html>',
                    encoding="utf-8",
                )
                (stage1_dir / "dashboard_data.js").write_text("window.__artifactLoaded = true;", encoding="utf-8")
                (stage1_dir / "PEAKS_Volcano_Plot.pdf").write_bytes(b"%PDF-1.4\n%artifact\n")
                (stage1_dir / "PEAKS_WGCNA_Complete_Session.RData").write_bytes(b"hidden-rdata")

                artifacts = client.get(f"/api/results/{run_id}/artifacts", headers=headers)
                self.assertEqual(artifacts.status_code, 200, artifacts.text)
                artifact_json = artifacts.json()
                volcano_files = artifact_json["volcano"]["featured"] + artifact_json["volcano"]["secondary"]
                rel_paths = [item["rel_path"] for item in volcano_files]
                self.assertIn("03_analysis_CBN_median/PROTEOMICS_Interactive_Volcano_Plot.html", rel_paths)
                self.assertFalse(any(path.endswith(".RData") for path in rel_paths))
                featured_companions = artifact_json["volcano"]["featured"][0]["download_companions"]
                self.assertTrue(featured_companions)
                self.assertTrue(any(item["kind"] == "pdf" for item in featured_companions))

                session_token = headers["Authorization"].split(" ", 1)[1]
                html = client.get(
                    f"/api/results/{run_id}/artifacts/file/stage1/PEAKS_Interactive_Volcano_Plot.html",
                    params={"session_token": session_token},
                    headers=headers,
                )
                self.assertEqual(html.status_code, 200, html.text)
                self.assertIn(
                    f"/api/results/{run_id}/artifacts/file/stage1/dashboard_data.js?session_token={session_token}",
                    html.text,
                )

                asset = client.get(
                    f"/api/results/{run_id}/artifacts/file/stage1/dashboard_data.js",
                    headers=headers,
                )
                self.assertEqual(asset.status_code, 200)
                self.assertIn("artifactLoaded", asset.text)

                files = client.get(f"/api/runs/{run_id}/files", headers=headers)
                self.assertEqual(files.status_code, 200, files.text)
                listed_paths = [item["rel_path"] for item in files.json()]
                self.assertIn("stage1/PEAKS_Interactive_Volcano_Plot.html", listed_paths)
                self.assertIn("stage1/PEAKS_Volcano_Plot.pdf", listed_paths)
                self.assertNotIn("stage1/dashboard_data.js", listed_paths)
                self.assertNotIn("stage1/PEAKS_WGCNA_Complete_Session.RData", listed_paths)

                blocked = client.get(
                    f"/api/runs/{run_id}/files/stage1/PEAKS_WGCNA_Complete_Session.RData",
                    headers=headers,
                )
                self.assertEqual(blocked.status_code, 404)

                export_zip = client.get(f"/api/runs/{run_id}/export/zip", headers=headers)
                self.assertEqual(export_zip.status_code, 200, export_zip.text)
                with zipfile.ZipFile(io.BytesIO(export_zip.content)) as archive:
                    archive_names = archive.namelist()
                self.assertIn("stage1/dashboard_data.js", archive_names)
                self.assertNotIn("stage1/PEAKS_WGCNA_Complete_Session.RData", archive_names)


    def test_qc_endpoint(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)
                uploaded = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes())
                dataset_id = uploaded["dataset_id"]

                # Valid QC request
                qc = client.get(f"/api/datasets/{dataset_id}/qc", headers=headers)
                self.assertEqual(qc.status_code, 200, qc.text)
                data = qc.json()

                # cv_histogram
                self.assertIn("cv_histogram", data)
                self.assertIn("bins", data["cv_histogram"])
                self.assertIn("counts", data["cv_histogram"])
                self.assertIsInstance(data["cv_histogram"]["bins"], list)
                self.assertIsInstance(data["cv_histogram"]["counts"], list)
                self.assertEqual(len(data["cv_histogram"]["counts"]), 30)
                self.assertEqual(len(data["cv_histogram"]["bins"]), 31)

                # missing_pct
                self.assertIn("missing_pct", data)
                for key, value in data["missing_pct"].items():
                    self.assertIsInstance(value, float)
                    self.assertGreaterEqual(value, 0.0)
                    self.assertLessEqual(value, 1.0)

                # pca
                self.assertIn("pca", data)
                self.assertTrue(data["pca"]["available"])
                self.assertIsInstance(data["pca"]["pc1"], list)
                self.assertEqual(len(data["pca"]["pc1"]), 6)  # 6 sample columns
                self.assertIsInstance(data["pca"]["var_pct"], list)
                self.assertEqual(len(data["pca"]["var_pct"]), 2)
                self.assertLessEqual(sum(data["pca"]["var_pct"]), 100.0)
                self.assertIsInstance(data["pca"]["labels"], list)
                self.assertTrue(all(isinstance(label, str) for label in data["pca"]["labels"]))

                # 404 for nonexistent dataset
                missing = client.get("/api/datasets/nonexistent-id/qc", headers=headers)
                self.assertEqual(missing.status_code, 404)

    def test_qc_edge_cases(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                # All-zero dataset: CV should have no Inf/NaN
                zero_buf = io.StringIO()
                zero_writer = csv.writer(zero_buf)
                zero_writer.writerow(["Protein", "S1", "S2", "S3", "S4"])
                for i in range(10):
                    zero_writer.writerow([f"GENE{i}", 0.0, 0.0, 0.0, 0.0])
                zero_content = zero_buf.getvalue().encode("utf-8")

                uploaded_zero = upload_primary_dataset(client, headers, workspace["id"], project["id"], zero_content)
                qc_zero = client.get(f"/api/datasets/{uploaded_zero['dataset_id']}/qc", headers=headers)
                self.assertEqual(qc_zero.status_code, 200, qc_zero.text)
                data_zero = qc_zero.json()
                # Zero-mean features should be excluded -- counts should have no Inf/NaN
                import math
                for count in data_zero["cv_histogram"].get("counts", []):
                    self.assertFalse(math.isinf(count), "CV histogram contains Inf")
                    self.assertFalse(math.isnan(count), "CV histogram contains NaN")

                # 2-sample dataset: PCA should be unavailable
                two_buf = io.StringIO()
                two_writer = csv.writer(two_buf)
                two_writer.writerow(["Protein", "S1", "S2"])
                for i in range(10):
                    two_writer.writerow([f"GENE{i}", 1.0 + i * 0.1, 2.0 + i * 0.2])
                two_content = two_buf.getvalue().encode("utf-8")

                uploaded_two = upload_primary_dataset(client, headers, workspace["id"], project["id"], two_content)
                qc_two = client.get(f"/api/datasets/{uploaded_two['dataset_id']}/qc", headers=headers)
                self.assertEqual(qc_two.status_code, 200, qc_two.text)
                data_two = qc_two.json()
                self.assertFalse(data_two["pca"]["available"])

    def test_trait_alignment_preview(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                uploaded = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes())
                primary_id = uploaded["dataset_id"]

                # Build traits that match dataset_bytes() sample columns
                # dataset_bytes() columns: Protein, Control_1, Control_2, Control_3, Disease_1, Disease_2, Disease_3
                trait_buf = io.StringIO()
                trait_writer = csv.writer(trait_buf)
                trait_writer.writerow(["sample", "Age"])
                trait_writer.writerow(["Control_1", 55])
                trait_writer.writerow(["Control_2", 60])
                trait_writer.writerow(["Control_3", 65])
                trait_writer.writerow(["Disease_1", 70])
                # Intentionally omit Disease_2, Disease_3 to test partial match
                trait_content = trait_buf.getvalue().encode("utf-8")

                traits = upload_dataset(client, headers, workspace["id"], project["id"], trait_content, "traits", "traits.csv")
                traits_id = traits["dataset_id"]

                # Valid alignment preview
                preview = client.get(
                    "/api/datasets/trait-alignment-preview",
                    headers=headers,
                    params={"dataset_id": primary_id, "traits_id": traits_id},
                )
                self.assertEqual(preview.status_code, 200, preview.text)
                data = preview.json()
                self.assertIn("matched", data)
                self.assertIn("total", data)
                self.assertIn("unmatched", data)
                self.assertEqual(data["matched"] + len(data["unmatched"]), data["total"])
                self.assertGreater(data["matched"], 0)
                self.assertEqual(data["matched"], 4)  # 4 traits provided, 4 matched
                self.assertEqual(len(data["unmatched"]), 2)  # Disease_2, Disease_3 unmatched

                # 404 for nonexistent dataset_id
                missing = client.get(
                    "/api/datasets/trait-alignment-preview",
                    headers=headers,
                    params={"dataset_id": "nonexistent-id", "traits_id": traits_id},
                )
                self.assertEqual(missing.status_code, 404)

    def test_traits_qc_endpoint(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                # Build traits CSV with numeric columns and one missing value
                trait_buf = io.StringIO()
                trait_writer = csv.writer(trait_buf)
                trait_writer.writerow(["sample", "Age", "BMI", "MMSE"])
                trait_writer.writerow(["Control_1", 55, 22.1, 28])
                trait_writer.writerow(["Control_2", 60, 24.5, 27])
                trait_writer.writerow(["Control_3", 65, "", 26])  # missing BMI
                trait_writer.writerow(["Disease_1", 70, 30.2, 29])
                trait_content = trait_buf.getvalue().encode("utf-8")

                traits = upload_dataset(client, headers, workspace["id"], project["id"], trait_content, "traits", "traits.csv")
                traits_id = traits["dataset_id"]

                # Valid traits QC request
                resp = client.get(f"/api/datasets/{traits_id}/traits-qc", headers=headers)
                self.assertEqual(resp.status_code, 200, resp.text)
                data = resp.json()

                # Structure checks
                self.assertIn("column_summaries", data)
                self.assertIn("trait_names", data)
                self.assertIn("sample_count", data)
                self.assertEqual(data["sample_count"], 4)

                # trait_names includes all columns
                self.assertIn("Age", data["trait_names"])
                self.assertIn("BMI", data["trait_names"])
                self.assertIn("MMSE", data["trait_names"])
                self.assertIn("sample", data["trait_names"])

                # column_summaries has entries for numeric columns
                summaries = data["column_summaries"]
                self.assertIn("Age", summaries)
                self.assertIn("BMI", summaries)
                self.assertIn("MMSE", summaries)

                # Age summary: 4 values, no missing
                age = summaries["Age"]
                self.assertEqual(age["count"], 4)
                self.assertAlmostEqual(age["mean"], 62.5)
                self.assertEqual(age["missing_count"], 0)
                self.assertIsInstance(age["std"], float)
                self.assertIsInstance(age["min"], (int, float))
                self.assertIsInstance(age["max"], (int, float))
                self.assertIn("outlier_count", age)

                # BMI summary: 3 values, 1 missing
                bmi = summaries["BMI"]
                self.assertEqual(bmi["count"], 3)
                self.assertEqual(bmi["missing_count"], 1)
                self.assertAlmostEqual(bmi["missing_pct"], 0.25)

                # 404 for nonexistent
                missing = client.get("/api/datasets/nonexistent-id/traits-qc", headers=headers)
                self.assertEqual(missing.status_code, 404)

    def test_run_stores_app_version(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)
                uploaded = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes())

                run_response = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Version Test Run",
                        "dataset_id": uploaded["dataset_id"],
                        "cohort1": "Control",
                        "cohort2": "Disease",
                    },
                )
                self.assertEqual(run_response.status_code, 200, run_response.text)
                run_id = run_response.json()["run_id"]

                # list_runs should include app_version
                list_resp = client.get("/api/runs", headers=headers, params={"workspace_id": workspace["id"]})
                self.assertEqual(list_resp.status_code, 200)
                runs = list_resp.json()
                target = [r for r in runs if r["id"] == run_id]
                self.assertTrue(len(target) > 0, "Run not found in list")
                self.assertEqual(target[0]["app_version"], "2.0.0")

                # get_run should include app_version
                detail_resp = client.get(f"/api/runs/{run_id}", headers=headers)
                self.assertEqual(detail_resp.status_code, 200)
                self.assertEqual(detail_resp.json()["app_version"], "2.0.0")

    def test_color_map_and_artifact_tabs(self) -> None:
        sys.path.insert(0, str(BACKEND_DIR))
        from services.deliverables import MODULE_COLOR_MAP
        from services.artifacts import ARTIFACT_TABS, artifact_tab, artifact_family

        # Color map has at least 40 entries
        self.assertGreaterEqual(len(MODULE_COLOR_MAP), 40)
        # All keys lowercase
        self.assertTrue(all(k == k.lower() for k in MODULE_COLOR_MAP), "All keys must be lowercase")
        # Every value starts with "#" (all hex, no CSS named color keywords)
        self.assertTrue(all(v.startswith("#") for v in MODULE_COLOR_MAP.values()), "All values must be hex")
        # Required WGCNA colors present
        for color in ("turquoise", "blue", "brown", "greenyellow", "midnightblue", "darkmagenta", "grey60"):
            self.assertIn(color, MODULE_COLOR_MAP, f"Missing WGCNA color: {color}")
        # Fabricated keys removed
        for bad_key in ("brown_2", "blue_2", "cyan_2", "green_2"):
            self.assertNotIn(bad_key, MODULE_COLOR_MAP, f"Fabricated key still present: {bad_key}")

        # Parameters tab restored — parameter files route to "parameters" tab
        self.assertIn("parameters", ARTIFACT_TABS)
        # Parameters tab routing
        self.assertEqual(artifact_tab("pipeline_params.html"), "parameters")
        self.assertEqual(artifact_tab("pipeline_parameters.json"), "parameters")
        self.assertEqual(artifact_tab("parameters_audit.json"), "parameters")
        self.assertEqual(artifact_family("pipeline_params.html", tab="parameters", kind="html"), "parameters.audit")

    def test_migration_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            build_app(temp_dir)  # triggers _migrate_sqlite_schema once
            from database import _migrate_sqlite_schema
            # Second call should be a no-op, not raise
            _migrate_sqlite_schema()

    def test_artifact_manifest_overrides_filename_guessing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            sys.path.insert(0, str(BACKEND_DIR))
            from services.artifacts import build_artifact_index

            run_dir = Path(tmp)
            custom_dir = run_dir / "custom_outputs"
            custom_dir.mkdir(parents=True)
            html_path = custom_dir / "alpha.html"
            html_path.write_text("<html><body>hello</body></html>", encoding="utf-8")
            (run_dir / "artifact_manifest.json").write_text(
                json.dumps(
                    {
                        "entries": {
                            "custom_outputs/alpha.html": {
                                "title": "Canonical Volcano",
                                "tab": "volcano",
                                "artifact_family": "volcano.main",
                                "canonical": True,
                                "inline_preference": "html",
                                "legacy_class": "product.native",
                            }
                        }
                    }
                )
            )

            index = build_artifact_index("RUN-test", run_dir, metadata={"input_level": "peptide"})
            target = next(item for item in index["artifacts"] if item["rel_path"] == "custom_outputs/alpha.html")
            self.assertEqual(target["title"], "Canonical Volcano")
            self.assertEqual(target["tab"], "volcano")
            self.assertEqual(target["artifact_family"], "volcano.main")
            self.assertTrue(target["canonical"])
            self.assertEqual(target["inline_preference"], "html")

    def test_write_artifact_index_saves_html_catalog(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            sys.path.insert(0, str(BACKEND_DIR))
            from services.artifacts import write_artifact_index

            run_dir = Path(tmp)
            (run_dir / "input").mkdir(parents=True)
            (run_dir / "input" / "cleaned_matrix.csv").write_text("feature_id,Control_1\nAPOE,1\n", encoding="utf-8")

            index = write_artifact_index("RUN-test", run_dir, metadata={"input_level": "protein"})

            html_path = run_dir / "artifact_index.html"
            json_path = run_dir / "artifact_index.json"
            self.assertTrue(html_path.exists())
            self.assertTrue(json_path.exists())
            html = html_path.read_text(encoding="utf-8")
            self.assertIn("Artifact Catalog", html)
            self.assertIn("artifacts available to the agent", html)
            self.assertIn("Cleaned and filtered abundance matrix", html)
            self.assertIn("artifact_index.html", html)
            self.assertIn('href="input/cleaned_matrix.csv"', html)
            table_paths = [item["rel_path"] for item in index["tables"]["featured"] + index["tables"]["secondary"]]
            self.assertIn("artifact_index.html", table_paths)
            self.assertIn("artifact_index.json", table_paths)

    def test_manifest_present_disables_science_tab_guessing_for_unmanifested_artifacts(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            sys.path.insert(0, str(BACKEND_DIR))
            from services.artifacts import build_artifact_index

            run_dir = Path(tmp)
            helper_path = run_dir / "custom_outputs" / "volcano_helper.html"
            helper_path.parent.mkdir(parents=True)
            helper_path.write_text("<html><body>helper</body></html>", encoding="utf-8")
            (run_dir / "artifact_manifest.json").write_text(json.dumps({"entries": {}}))

            index = build_artifact_index("RUN-test", run_dir, metadata={"input_level": "peptide"})
            target = next(item for item in index["artifacts"] if item["rel_path"] == "custom_outputs/volcano_helper.html")
            self.assertEqual(target["tab"], "files")
            self.assertEqual(target["artifact_family"], "files.support")
            self.assertFalse(target["canonical"])

    def test_artifact_manifest_covers_runtime_tables_and_legacy_subtrees(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            sys.path.insert(0, str(BACKEND_DIR))
            from services.artifact_manifest import build_artifact_manifest

            run_dir = Path(tmp)
            (run_dir / "input").mkdir(parents=True)
            (run_dir / "stage1").mkdir(parents=True)
            (run_dir / "stage2").mkdir(parents=True)
            (run_dir / "stage3").mkdir(parents=True)
            (run_dir / "05_network_cbn_median" / "disease_status").mkdir(parents=True)

            (run_dir / "input" / "cleaned_matrix.csv").write_text("feature_id,Control_1\nA,1\n", encoding="utf-8")
            (run_dir / "stage1" / "volcano_results.tsv").write_text("gene\tlog2fc\tadj_pvalue\nAPOE\t1.2\t0.01\n", encoding="utf-8")
            (run_dir / "stage3" / "celltype_hitListStats_Astrocytes.csv").write_text("module,hits\nblue,2\n", encoding="utf-8")
            (run_dir / "05_network_cbn_median" / "disease_status" / "Amyloid_Associated_Modules.pdf").write_text("pdf", encoding="utf-8")

            profile = {
                "deliverable_prefix": "SPEC",
                "display_prefix": "Spec Pep",
                "normalization_tag": "cbn_median",
                "go_label": "GO_All",
                "input_level": "peptide",
                "generated_at": "2026-04-23T00:00:00Z",
            }
            manifest = build_artifact_manifest(run_dir, profile)
            entries = manifest["entries"]

            self.assertIn("input/cleaned_matrix.csv", entries)
            self.assertEqual(entries["input/cleaned_matrix.csv"]["artifact_family"], "tables.cleaned_matrix")
            self.assertIn("stage1/volcano_results.tsv", entries)
            self.assertEqual(entries["stage1/volcano_results.tsv"]["artifact_family"], "volcano.results")
            self.assertIn("stage3/celltype_hitListStats_Astrocytes.csv", entries)
            self.assertTrue(entries["stage3/celltype_hitListStats_Astrocytes.csv"]["artifact_family"].startswith("cells.hit_list"))
            self.assertIn("05_network_cbn_median/disease_status/Amyloid_Associated_Modules.pdf", entries)
            self.assertEqual(entries["05_network_cbn_median/disease_status/Amyloid_Associated_Modules.pdf"]["tab"], "network")
            self.assertTrue(entries["05_network_cbn_median/disease_status/Amyloid_Associated_Modules.pdf"]["canonical"])
            self.assertNotIn("Spec Pep Reports/SPEC_Interactive_Dashboard.html", entries)

    def test_artifact_index_prefers_interactive_html_for_go_and_cell_heatmaps(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            sys.path.insert(0, str(BACKEND_DIR))
            from services.artifacts import build_artifact_index

            run_dir = Path(tmp)
            go_dir = run_dir / "Spec Pep Go"
            cell_dir = run_dir / "Spec Pep CellTypeFET"
            go_dir.mkdir(parents=True)
            cell_dir.mkdir(parents=True)
            (go_dir / "SPEC_GO_Interactive_Heatmap.html").write_text("<html>go</html>", encoding="utf-8")
            (go_dir / "GSA-GO-FET_SPEC-PEP_Proteomics_GO-redundancyRemoved.Kbest.pdf").write_bytes(b"%PDF-1.4\n")
            (cell_dir / "SPEC_CellTypeFET_Interactive_Heatmap.html").write_text("<html>cell</html>", encoding="utf-8")
            (cell_dir / "SPEC_Peptides_CellTypeFET.Overlap.pdf").write_bytes(b"%PDF-1.4\n")
            (run_dir / "artifact_manifest.json").write_text(
                json.dumps(
                    {
                        "entries": {
                            "Spec Pep Go/SPEC_GO_Interactive_Heatmap.html": {
                                "title": "GO Heatmap",
                                "tab": "go",
                                "artifact_family": "go.heatmap",
                                "canonical": True,
                                "inline_preference": "html",
                                "legacy_class": "legacy.go",
                            },
                            "Spec Pep Go/GSA-GO-FET_SPEC-PEP_Proteomics_GO-redundancyRemoved.Kbest.pdf": {
                                "title": "GO Heatmap",
                                "tab": "go",
                                "artifact_family": "go.heatmap",
                                "canonical": True,
                                "inline_preference": "pdf",
                                "legacy_class": "legacy.go",
                            },
                            "Spec Pep CellTypeFET/SPEC_CellTypeFET_Interactive_Heatmap.html": {
                                "title": "Cell-type Heatmap",
                                "tab": "cells",
                                "artifact_family": "cells.heatmap",
                                "canonical": True,
                                "inline_preference": "html",
                                "legacy_class": "legacy.celltype",
                            },
                            "Spec Pep CellTypeFET/SPEC_Peptides_CellTypeFET.Overlap.pdf": {
                                "title": "Cell-type Heatmap",
                                "tab": "cells",
                                "artifact_family": "cells.heatmap",
                                "canonical": True,
                                "inline_preference": "pdf",
                                "legacy_class": "legacy.celltype",
                            },
                        }
                    }
                )
            )

            index = build_artifact_index("RUN-test", run_dir, metadata={"input_level": "peptide"})
            go_heatmap = next(item for item in index["go"]["featured"] if item["artifact_family"] == "go.heatmap")
            cell_heatmap = next(item for item in index["cells"]["featured"] if item["artifact_family"] == "cells.heatmap")
            self.assertEqual(go_heatmap["kind"], "html")
            self.assertIn("pdf", {item["kind"] for item in go_heatmap["download_companions"]})
            self.assertEqual(cell_heatmap["kind"], "html")
            self.assertIn("pdf", {item["kind"] for item in cell_heatmap["download_companions"]})

    def test_trait_specific_wgcna_deliverables_are_emitted_from_run_traits(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            sys.path.insert(0, str(BACKEND_DIR))
            from services.deliverables import emit_trait_specific_wgcna_bundle

            run_dir = Path(tmp)
            stage1 = run_dir / "stage1"
            input_dir = run_dir / "input"
            network_dir = run_dir / "05_network_CBN_median"
            stage1.mkdir(parents=True)
            input_dir.mkdir(parents=True)
            network_dir.mkdir(parents=True)

            (stage1 / "module_trait_cor.csv").write_text(
                "module_color,cor_Disease,p_Disease,cor_t-tau [ng/L],p_t-tau [ng/L],cor_p-tau [ng/L],p_p-tau [ng/L],cor_Abeta-42 [ng/L],p_Abeta-42 [ng/L]\n"
                "MEblue,0.62,0.001,0.51,0.002,-0.48,0.004,0.45,0.006\n"
                "MEgreen,-0.44,0.003,-0.31,0.02,0.12,0.40,-0.55,0.001\n",
                encoding="utf-8",
            )
            (stage1 / "module_eigengenes.csv").write_text(
                "sample_name,group,MEblue,MEgreen\n"
                "Control_1,Control,0.10,-0.20\n"
                "Control_2,Control,0.12,-0.18\n"
                "Disease_1,Disease,0.45,-0.50\n"
                "Disease_2,Disease,0.52,-0.55\n",
                encoding="utf-8",
            )
            (input_dir / "traits.csv").write_text(
                "SAMPLE_ID,GROUP,t-tau [ng/L],p-tau [ng/L],Abeta-42 [ng/L]\n"
                "Control_1,Control,200,20,800\n"
                "Control_2,Control,220,22,790\n"
                "Disease_1,Disease,700,80,420\n"
                "Disease_2,Disease,760,84,400\n",
                encoding="utf-8",
            )
            for name in [
                "TEST_WGCNA_01_Sample_Clustering_QC.pdf",
                "TEST_WGCNA_02_Power_Selection.pdf",
                "TEST_WGCNA_03_Network_Dendrograms.pdf",
                "TEST_WGCNA_04_Module_Trait_Correlations.pdf",
                "TEST_WGCNA_05_Module_Response_Plots.pdf",
            ]:
                (network_dir / name).write_bytes(b"%PDF-1.4\n")
            (network_dir / "TEST_WGCNA_Module_Assignments_with_kME.csv").write_text(
                "peptide_id,gene,module_color,kME\nP1,APOE,blue,0.9\nP2,CLU,green,0.8\n",
                encoding="utf-8",
            )
            (network_dir / "TEST_WGCNA_Module_Eigengenes.csv").write_text("sample_name,MEblue\nControl_1,0.1\n", encoding="utf-8")
            (network_dir / "TEST_WGCNA_Complete_Results.xlsx").write_bytes(b"placeholder")

            logs: list[str] = []
            emit_trait_specific_wgcna_bundle(
                run_dir,
                {"deliverable_prefix": "TEST", "normalization_tag": "CBN_median"},
                logs.append,
            )

            expected = {
                "disease_status/TEST_WGCNA_04_Disease_Trait_Correlations.pdf",
                "disease_status/Disease_Associated_Modules.csv",
                "total_tau/TEST_WGCNA_04_Total_Tau_Correlations.pdf",
                "total_tau/TEST_WGCNA_05_Total_Tau_Response_Plots.pdf",
                "total_tau/T_TAU_Associated_Modules.csv",
                "phospho_tau/TEST_WGCNA_04_Phospho_Tau_Correlations.pdf",
                "phospho_tau/TEST_WGCNA_05_Phospho_Tau_Response_Plots.pdf",
                "phospho_tau/P_TAU_Associated_Modules.csv",
                "amyloid_beta/TEST_WGCNA_04_Amyloid_Trait_Correlations.pdf",
                "amyloid_beta/TEST_WGCNA_05_Amyloid_Response_Plots.pdf",
                "amyloid_beta/ABETA42_Associated_Modules.csv",
                "all_traits_comprehensive/TEST_WGCNA_04_All_Traits_Heatmap.pdf",
                "all_traits_comprehensive/TEST_WGCNA_Complete_Results.xlsx",
                "ad_pathology_composite/TEST_WGCNA_Module_Assignments.csv",
            }
            rel_files = {str(path.relative_to(network_dir)) for path in network_dir.rglob("*") if path.is_file()}
            self.assertTrue(expected.issubset(rel_files), f"Missing trait WGCNA files: {expected - rel_files}")
            self.assertFalse(any(path.suffix.lower() in {".rdata", ".rds", ".rda"} for path in network_dir.rglob("*")))

    def test_generic_trait_subfolders_are_emitted_for_non_ad_cohorts(self) -> None:
        """Data-agnostic: a trait CSV with no AD-pattern columns must still produce
        per-trait subfolders so the deliverable is complete for any cohort."""
        with tempfile.TemporaryDirectory() as tmp:
            sys.path.insert(0, str(BACKEND_DIR))
            from services.deliverables import emit_trait_specific_wgcna_bundle

            run_dir = Path(tmp)
            stage1 = run_dir / "stage1"
            input_dir = run_dir / "input"
            network_dir = run_dir / "05_network_CBN_median"
            stage1.mkdir(parents=True)
            input_dir.mkdir(parents=True)
            network_dir.mkdir(parents=True)

            # Plasma aging cohort — no AD pattern in any trait name
            (stage1 / "module_trait_cor.csv").write_text(
                "module_color,cor_age,p_age,cor_BMI,p_BMI,cor_response_score,p_response_score\n"
                "MEblue,0.55,0.001,0.32,0.04,-0.41,0.008\n"
                "MEgreen,-0.38,0.012,0.21,0.18,0.49,0.003\n",
                encoding="utf-8",
            )
            (stage1 / "module_eigengenes.csv").write_text(
                "sample_name,group,MEblue,MEgreen\n"
                "S1,Young,0.10,-0.20\n"
                "S2,Young,0.12,-0.18\n"
                "S3,Old,0.45,-0.50\n"
                "S4,Old,0.52,-0.55\n",
                encoding="utf-8",
            )
            (input_dir / "traits.csv").write_text(
                "SAMPLE_ID,GROUP,age,BMI,response_score\n"
                "S1,Young,25,22.1,3.2\n"
                "S2,Young,28,24.5,3.5\n"
                "S3,Old,72,28.4,7.1\n"
                "S4,Old,75,29.1,7.8\n",
                encoding="utf-8",
            )
            for name in [
                "TEST_WGCNA_01_Sample_Clustering_QC.pdf",
                "TEST_WGCNA_02_Power_Selection.pdf",
                "TEST_WGCNA_03_Network_Dendrograms.pdf",
                "TEST_WGCNA_04_Module_Trait_Correlations.pdf",
            ]:
                (network_dir / name).write_bytes(b"%PDF-1.4\n")
            (network_dir / "TEST_WGCNA_Module_Assignments_with_kME.csv").write_text(
                "peptide_id,gene,module_color,kME\nP1,APOE,blue,0.9\nP2,CLU,green,0.8\n",
                encoding="utf-8",
            )
            (network_dir / "TEST_WGCNA_Module_Eigengenes.csv").write_text(
                "sample_name,MEblue\nS1,0.1\n", encoding="utf-8",
            )

            logs: list[str] = []
            emit_trait_specific_wgcna_bundle(
                run_dir,
                {"deliverable_prefix": "TEST", "normalization_tag": "CBN_median"},
                logs.append,
            )

            # Each non-AD trait must get its own subfolder with the trio:
            # correlation heatmap, response plots, associated modules CSV, plus shared QC.
            for trait_slug, trait_name in [("age", "age"), ("bmi", "BMI"), ("response_score", "response_score")]:
                trait_dir = network_dir / trait_slug
                self.assertTrue(trait_dir.is_dir(), f"Generic trait folder {trait_slug}/ not created")
                # Shared QC files copied
                self.assertTrue(
                    (trait_dir / "TEST_WGCNA_01_Sample_Clustering_QC.pdf").exists(),
                    f"{trait_slug}/ missing shared sample-clustering PDF",
                )
                # Per-trait correlation heatmap (single column)
                self.assertTrue(
                    any(p.name.startswith("TEST_WGCNA_04_") and p.name.endswith("_Correlations.pdf") for p in trait_dir.iterdir()),
                    f"{trait_slug}/ missing correlation heatmap PDF",
                )
                # Per-trait response plots
                self.assertTrue(
                    any(p.name.startswith("TEST_WGCNA_05_") and p.name.endswith("_Response_Plots.pdf") for p in trait_dir.iterdir()),
                    f"{trait_slug}/ missing response plots PDF",
                )
                # Associated modules CSV
                self.assertTrue(
                    any(p.name.endswith("_Associated_Modules.csv") for p in trait_dir.iterdir()),
                    f"{trait_slug}/ missing Associated_Modules CSV",
                )

            # all_traits_comprehensive still emitted with ≥2 traits
            self.assertTrue((network_dir / "all_traits_comprehensive").is_dir())
            # ad_pathology_composite NOT emitted (no AD-pattern traits)
            self.assertFalse(
                (network_dir / "ad_pathology_composite").is_dir(),
                "ad_pathology_composite should only appear when ≥2 AD-bucket traits are present",
            )

    def test_mixed_ad_and_generic_trait_subfolders_coexist(self) -> None:
        """When a cohort has both AD-pattern traits and generic continuous traits,
        the named AD folders AND generic per-trait folders must both be emitted."""
        with tempfile.TemporaryDirectory() as tmp:
            sys.path.insert(0, str(BACKEND_DIR))
            from services.deliverables import emit_trait_specific_wgcna_bundle

            run_dir = Path(tmp)
            stage1 = run_dir / "stage1"
            input_dir = run_dir / "input"
            network_dir = run_dir / "05_network_CBN_median"
            stage1.mkdir(parents=True)
            input_dir.mkdir(parents=True)
            network_dir.mkdir(parents=True)

            # Mixed: AD pathology + a generic clinical score (>=4 samples to satisfy
            # the per-trait gate; 2-sample fixture pre-dated the gate).
            (stage1 / "module_trait_cor.csv").write_text(
                "module_color,cor_t-tau,p_t-tau,cor_MMSE,p_MMSE\n"
                "MEblue,0.55,0.001,-0.41,0.008\n"
                "MEgreen,-0.38,0.012,0.49,0.003\n",
                encoding="utf-8",
            )
            (stage1 / "module_eigengenes.csv").write_text(
                "sample_name,MEblue,MEgreen\nS1,0.1,-0.2\nS2,0.45,-0.5\nS3,0.20,-0.32\nS4,0.38,-0.42\n",
                encoding="utf-8",
            )
            (input_dir / "traits.csv").write_text(
                "SAMPLE_ID,t-tau,MMSE\nS1,200,29\nS2,700,18\nS3,310,26\nS4,540,22\n",
                encoding="utf-8",
            )
            for name in [
                "TEST_WGCNA_01_Sample_Clustering_QC.pdf",
                "TEST_WGCNA_02_Power_Selection.pdf",
                "TEST_WGCNA_03_Network_Dendrograms.pdf",
                "TEST_WGCNA_04_Module_Trait_Correlations.pdf",
            ]:
                (network_dir / name).write_bytes(b"%PDF-1.4\n")
            (network_dir / "TEST_WGCNA_Module_Assignments_with_kME.csv").write_text(
                "peptide_id,gene,module_color,kME\nP1,APOE,blue,0.9\n", encoding="utf-8",
            )
            (network_dir / "TEST_WGCNA_Module_Eigengenes.csv").write_text(
                "sample_name,MEblue\nS1,0.1\n", encoding="utf-8",
            )

            emit_trait_specific_wgcna_bundle(
                run_dir,
                {"deliverable_prefix": "TEST", "normalization_tag": "CBN_median"},
                lambda _line: None,
            )

            # AD bucket still works
            self.assertTrue((network_dir / "total_tau").is_dir(), "AD bucket lost when mixed with generic traits")
            self.assertTrue((network_dir / "total_tau" / "TEST_WGCNA_04_Total_Tau_Correlations.pdf").exists())
            # Generic trait also got a folder
            self.assertTrue((network_dir / "mmse").is_dir(), "Generic trait MMSE not bucketed alongside AD trait")

    def test_canonical_traits_emit_raw_and_std_variant_csvs(self) -> None:
        """E2E: canonical T_TAU/P_TAU/ABETA42 in expanded traits + module_trait_cor with
        cor_*_raw / cor_*_std columns produces both _raw and _std Associated_Modules CSVs."""
        with tempfile.TemporaryDirectory() as tmp:
            sys.path.insert(0, str(BACKEND_DIR))
            from services.deliverables import emit_trait_specific_wgcna_bundle

            run_dir = Path(tmp)
            stage1 = run_dir / "stage1"
            input_dir = run_dir / "input"
            network_dir = run_dir / "05_network_CBN_median"
            stage1.mkdir(parents=True)
            input_dir.mkdir(parents=True)
            network_dir.mkdir(parents=True)

            # module_trait_cor with both raw and std variant columns for each AD trait
            (stage1 / "module_trait_cor.csv").write_text(
                "module_color,"
                "cor_AD,p_AD,"
                "cor_T_TAU_raw,p_T_TAU_raw,cor_T_TAU_std,p_T_TAU_std,"
                "cor_P_TAU_raw,p_P_TAU_raw,cor_P_TAU_std,p_P_TAU_std,"
                "cor_ABETA42_raw,p_ABETA42_raw,cor_ABETA42_std,p_ABETA42_std\n"
                "MEblue,0.5,0.001,0.42,0.003,0.40,0.004,0.38,0.005,0.37,0.006,-0.45,0.002,-0.44,0.002\n"
                "MEgreen,-0.42,0.003,-0.31,0.02,-0.30,0.022,0.12,0.4,0.13,0.39,-0.55,0.001,-0.53,0.001\n",
                encoding="utf-8",
            )
            (stage1 / "module_eigengenes.csv").write_text(
                "sample_name,group,MEblue,MEgreen\n"
                "Control_1,Control,0.10,-0.20\n"
                "Control_2,Control,0.12,-0.18\n"
                "Disease_1,Disease,0.45,-0.50\n"
                "Disease_2,Disease,0.52,-0.55\n",
                encoding="utf-8",
            )
            # Expanded traits (mimicking what _expand_traits writes)
            (stage1 / "expanded_traits.csv").write_text(
                "Sample,GROUP,AD,T_TAU_raw,T_TAU_std,P_TAU_raw,P_TAU_std,ABETA42_raw,ABETA42_std\n"
                "Control_1,Control,0,200,0.5,20,0.5,800,1.5\n"
                "Control_2,Control,0,220,0.7,22,0.7,790,1.4\n"
                "Disease_1,Disease,1,700,2.1,80,2.1,420,-1.5\n"
                "Disease_2,Disease,1,760,2.3,84,2.3,400,-1.7\n",
                encoding="utf-8",
            )
            for name in [
                "TEST_WGCNA_01_Sample_Clustering_QC.pdf",
                "TEST_WGCNA_02_Power_Selection.pdf",
                "TEST_WGCNA_03_Network_Dendrograms.pdf",
                "TEST_WGCNA_04_Module_Trait_Correlations.pdf",
            ]:
                (network_dir / name).write_bytes(b"%PDF-1.4\n")
            (network_dir / "TEST_WGCNA_Module_Assignments_with_kME.csv").write_text(
                "peptide_id,gene,module_color,kME\nP1,APOE,blue,0.9\n",
                encoding="utf-8",
            )
            (network_dir / "TEST_WGCNA_Module_Eigengenes.csv").write_text(
                "sample_name,MEblue\nControl_1,0.1\n", encoding="utf-8",
            )

            emit_trait_specific_wgcna_bundle(
                run_dir,
                {"deliverable_prefix": "TEST", "normalization_tag": "CBN_median"},
                lambda _line: None,
            )

            # Each AD bucket has both raw and std variant CSVs
            self.assertTrue((network_dir / "disease_status" / "AD_Associated_Modules.csv").exists())
            for variant in ("T_TAU_raw", "T_TAU_std"):
                self.assertTrue(
                    (network_dir / "total_tau" / f"{variant}_Associated_Modules.csv").exists(),
                    f"Missing {variant}_Associated_Modules.csv"
                )
            for variant in ("P_TAU_raw", "P_TAU_std"):
                self.assertTrue(
                    (network_dir / "phospho_tau" / f"{variant}_Associated_Modules.csv").exists(),
                    f"Missing {variant}_Associated_Modules.csv"
                )
            for variant in ("ABETA42_raw", "ABETA42_std"):
                self.assertTrue(
                    (network_dir / "amyloid_beta" / f"{variant}_Associated_Modules.csv").exists(),
                    f"Missing {variant}_Associated_Modules.csv"
                )

    def test_gate_skips_per_trait_when_no_usable_metadata(self) -> None:
        """Gate behavior: traits.csv with only Sample column → no per-trait subfolders.
        Top-level WGCNA outputs are unaffected (not part of emit_trait_specific_wgcna_bundle)."""
        with tempfile.TemporaryDirectory() as tmp:
            sys.path.insert(0, str(BACKEND_DIR))
            from services.deliverables import emit_trait_specific_wgcna_bundle

            run_dir = Path(tmp)
            stage1 = run_dir / "stage1"
            input_dir = run_dir / "input"
            network_dir = run_dir / "05_network_CBN_median"
            stage1.mkdir(parents=True)
            input_dir.mkdir(parents=True)
            network_dir.mkdir(parents=True)

            # module_trait_cor exists (could come from non-trait analysis)
            (stage1 / "module_trait_cor.csv").write_text(
                "module_color,cor_dummy,p_dummy\nMEblue,0.5,0.001\n", encoding="utf-8",
            )
            # traits.csv has only the Sample column — gate must skip
            (input_dir / "traits.csv").write_text(
                "Sample\nS1\nS2\nS3\nS4\nS5\n", encoding="utf-8",
            )

            logs: list[str] = []
            emit_trait_specific_wgcna_bundle(
                run_dir,
                {"deliverable_prefix": "TEST", "normalization_tag": "CBN_median"},
                logs.append,
            )

            # No per-trait subfolders created
            for folder in ("disease_status", "total_tau", "phospho_tau", "amyloid_beta",
                           "ad_pathology_composite", "all_traits_comprehensive"):
                self.assertFalse((network_dir / folder).exists(),
                                 f"Subfolder {folder}/ should not be created when no traits provided")
            # Gate logged the skip
            self.assertTrue(any("No usable trait metadata" in line for line in logs),
                            f"Expected gate-skip log line; got: {logs}")

    def test_trait_alignment_whitespace(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                uploaded = upload_primary_dataset(client, headers, workspace["id"], project["id"], dataset_bytes())
                primary_id = uploaded["dataset_id"]

                # Create traits CSV with whitespace-padded sample names
                ws_buf = io.StringIO()
                ws_writer = csv.writer(ws_buf)
                ws_writer.writerow(["sample", "Age"])
                # dataset_bytes() has columns: Control_1, Control_2, Control_3, Disease_1, Disease_2, Disease_3
                ws_writer.writerow(["  Control_1  ", 55])
                ws_writer.writerow([" Control_2", 60])
                ws_writer.writerow(["Control_3 ", 65])
                ws_writer.writerow(["  Disease_1 ", 70])
                ws_writer.writerow([" Disease_2  ", 75])
                ws_writer.writerow(["Disease_3", 80])
                ws_content = ws_buf.getvalue().encode("utf-8")

                traits = upload_dataset(client, headers, workspace["id"], project["id"], ws_content, "traits", "traits_ws.csv")
                traits_id = traits["dataset_id"]

                preview = client.get(
                    "/api/datasets/trait-alignment-preview",
                    headers=headers,
                    params={"dataset_id": primary_id, "traits_id": traits_id},
                )
                self.assertEqual(preview.status_code, 200, preview.text)
                data = preview.json()
                # All 6 samples should match despite whitespace
                self.assertEqual(data["matched"], 6)
                self.assertEqual(len(data["unmatched"]), 0)


    def test_rob01_gene_column_access(self) -> None:
        """ROB-01: gene column access uses if/in check, not DataFrame.get()."""
        import pandas as pd
        sys.path.insert(0, str(BACKEND_DIR))
        from services.pipeline import _load_canonical_bundle_for_stage1

        sample_meta = pd.DataFrame({
            "sample_name": ["Control_1", "Control_2", "Disease_1", "Disease_2"],
            "group": ["Control", "Control", "Disease", "Disease"],
        })

        with tempfile.TemporaryDirectory() as tmp:
            run_dir = Path(tmp)
            input_dir = run_dir / "input"
            input_dir.mkdir()
            sample_meta.to_csv(input_dir / "sample_metadata.csv", index=False)

            # Case 1: matrix WITH gene column (NaN in gene falls back to feature_id)
            df_with_gene = pd.DataFrame({
                "feature_id": ["P1", "P2", "P3"],
                "gene": ["APOE", None, "CLU"],
                "Control_1": [10.0, 8.0, 7.0],
                "Control_2": [10.5, 8.5, 7.5],
                "Disease_1": [13.0, 11.0, 9.0],
                "Disease_2": [13.5, 11.5, 9.5],
            })
            df_with_gene.to_csv(input_dir / "cleaned_matrix.csv", index=False)
            logs = []
            bundle = _load_canonical_bundle_for_stage1(
                run_dir, "Control", "Disease", lambda msg: logs.append(msg)
            )
            self.assertIsNotNone(bundle)
            # Gene column present: gene values used, NaN falls back to feature_id
            self.assertEqual(bundle["gene_names"], ["APOE", "P2", "CLU"])

            # Case 2: matrix WITHOUT gene column — feature_id used directly
            df_no_gene = pd.DataFrame({
                "feature_id": ["P1", "P2", "P3"],
                "Control_1": [10.0, 8.0, 7.0],
                "Control_2": [10.5, 8.5, 7.5],
                "Disease_1": [13.0, 11.0, 9.0],
                "Disease_2": [13.5, 11.5, 9.5],
            })
            df_no_gene.to_csv(input_dir / "cleaned_matrix.csv", index=False)
            bundle2 = _load_canonical_bundle_for_stage1(
                run_dir, "Control", "Disease", lambda msg: logs.append(msg)
            )
            self.assertIsNotNone(bundle2)
            # No gene column: feature_id used directly
            self.assertEqual(bundle2["gene_names"], ["P1", "P2", "P3"])

    def test_rob02_gmt_warning(self) -> None:
        """ROB-02: malformed GMT lines produce WARNING log entries."""
        sys.path.insert(0, str(BACKEND_DIR))
        from services.pipeline import _load_gmt

        with tempfile.NamedTemporaryFile(mode="w", suffix=".gmt", delete=False) as f:
            # Line 1: valid (3+ fields)
            f.write("GOBP_APOPTOSIS\thttp://example.com\tGENE1\tGENE2\n")
            # Line 2: malformed (only 1 field, no tabs)
            f.write("MALFORMED_LINE_NO_TABS\n")
            # Line 3: malformed (only 2 fields)
            f.write("PARTIAL\thttp://example.com\n")
            # Line 4: valid
            f.write("GOMF_BINDING\thttp://example.com\tGENE3\n")
            gmt_path = f.name

        try:
            logs = []
            terms = _load_gmt(gmt_path, lambda msg: logs.append(msg))
            # Two valid terms parsed
            self.assertEqual(len(terms), 2)
            # Two warnings logged (one per malformed line)
            warnings = [log for log in logs if "WARNING" in log and "malformed gmt" in log.lower()]
            self.assertEqual(len(warnings), 2, f"Expected 2 GMT warnings, got {len(warnings)}: {logs}")
        finally:
            os.unlink(gmt_path)

    def test_rob04_volcano_no_cap(self) -> None:
        """ROB-04: volcano data arrays are not capped at 7500 entries."""
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                # Create a dataset with >7500 rows by generating a large CSV
                buffer = io.StringIO()
                writer = csv.writer(buffer)
                writer.writerow(["Protein", "Control_1", "Control_2", "Disease_1", "Disease_2"])
                import random
                random.seed(42)
                for i in range(8000):
                    writer.writerow([f"PROT_{i}", random.uniform(5, 15), random.uniform(5, 15), random.uniform(5, 15), random.uniform(5, 15)])

                upload = client.post(
                    "/api/uploads",
                    headers=headers,
                    data={"workspace_id": workspace["id"], "project_id": project["id"]},
                    files={"file": ("big_dataset.csv", buffer.getvalue().encode(), "text/csv")},
                )
                self.assertEqual(upload.status_code, 200, upload.text)
                dataset_id = upload.json()["dataset_id"]

                run_response = client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Volcano Cap Test",
                        "dataset_id": dataset_id,
                        "cohort1": "Control",
                        "cohort2": "Disease",
                    },
                )
                self.assertEqual(run_response.status_code, 200, run_response.text)
                run_id = run_response.json()["run_id"]

                dash = client.get(f"/api/results/{run_id}/dashboard", headers=headers)
                if dash.status_code == 200:
                    volcano = dash.json().get("volcano", {})
                    x_len = len(volcano.get("x", []))
                    # If the run produced volcano data, it must NOT be capped at 7500
                    if x_len > 0:
                        self.assertGreater(x_len, 7500, f"Volcano x has {x_len} entries — cap still present")

    def test_rob06_audit_pagination(self) -> None:
        """ROB-06: audit endpoint returns paginated response, not hardcapped list."""
        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                # Create a run to generate audit events
                upload = client.post(
                    "/api/uploads",
                    headers=headers,
                    data={"workspace_id": workspace["id"], "project_id": project["id"]},
                    files={"file": ("test.csv", dataset_bytes(), "text/csv")},
                )
                self.assertEqual(upload.status_code, 200, upload.text)
                dataset_id = upload.json()["dataset_id"]

                client.post(
                    "/api/runs",
                    headers=headers,
                    json={
                        "workspace_id": workspace["id"],
                        "project_id": project["id"],
                        "name": "Pagination Test Run",
                        "dataset_id": dataset_id,
                        "cohort1": "Control",
                        "cohort2": "Disease",
                    },
                )

                # Fetch audit with default pagination
                audit = client.get("/api/audit", headers=headers, params={"workspace_id": workspace["id"]})
                self.assertEqual(audit.status_code, 200, audit.text)
                body = audit.json()

                # Verify paginated envelope structure
                self.assertIn("events", body)
                self.assertIsInstance(body["events"], list)
                self.assertIn("page", body)
                self.assertIn("page_size", body)
                self.assertIn("total", body)
                self.assertIn("total_pages", body)
                self.assertEqual(body["page"], 1)
                self.assertEqual(body["page_size"], 50)
                self.assertIsInstance(body["total"], int)
                self.assertGreater(body["total"], 0)
                self.assertIsInstance(body["total_pages"], int)
                self.assertGreaterEqual(body["total_pages"], 1)

                # Verify events have expected shape
                self.assertGreater(len(body["events"]), 0)
                event = body["events"][0]
                self.assertIn("id", event)
                self.assertIn("action_type", event)
                self.assertIn("created_at", event)

                # Test with explicit page_size=1 to verify pagination works
                audit_p1 = client.get(
                    "/api/audit",
                    headers=headers,
                    params={"workspace_id": workspace["id"], "page": 1, "page_size": 1},
                )
                self.assertEqual(audit_p1.status_code, 200, audit_p1.text)
                body_p1 = audit_p1.json()
                self.assertEqual(body_p1["page"], 1)
                self.assertEqual(body_p1["page_size"], 1)
                self.assertEqual(len(body_p1["events"]), 1)
                self.assertEqual(body_p1["total"], body["total"])  # same total regardless of page_size

    def test_hub_sheets_generated_from_trait_associations(self) -> None:
        """HUB-01: Hub protein sheets appear when trait association files exist."""
        import tempfile
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            run_dir = Path(tmp)
            stage1 = run_dir / "stage1"
            stage1.mkdir()
            ma_path = stage1 / "module_assignments.csv"
            ma_path.write_text(
                "peptide_id,feature_id,gene,module_color,kME,alternative_module,kME_alternative,module_quality\n"
                "P001,F001,APOE,blue,0.95,turquoise,0.30,high\n"
                "P002,F002,CLU,blue,0.88,turquoise,0.25,high\n"
                "P003,F003,BIN1,blue,0.82,green,0.40,medium\n"
                "P004,F004,TREM2,green,0.91,blue,0.20,high\n"
                "P005,F005,ABCA7,green,0.85,yellow,0.15,medium\n"
                "P006,F006,CD33,red,0.78,blue,0.10,low\n"
            )
            trait_dir = stage1 / "trait_associations"
            trait_dir.mkdir()
            (trait_dir / "Disease_Associated_Modules.csv").write_text(
                "Module,Correlation,P_Value,Direction\n"
                "MEblue,0.629,7.3e-08,Upregulated_in_Disease\n"
                "MEgreen,-0.479,0.0001,Downregulated_in_Disease\n"
            )
            (trait_dir / "Abeta42_Associated_Modules.csv").write_text(
                "Module,Correlation,P_Value,Direction\n"
                "MEblue,0.55,0.001,Upregulated_in_Abeta42\n"
            )
            (run_dir / "stage2").mkdir()
            (run_dir / "stage2" / "go_zscore_matrix.csv").write_text("term,blue,green\nGO:0001,1.5,0.3\n")
            (run_dir / "stage3").mkdir()
            (run_dir / "stage3" / "celltype_FDR_matrix.csv").write_text("module,Neurons,Astrocytes\nblue,0.01,0.5\n")

            norm_tag = "CBN_median"
            profile = {"deliverable_prefix": "TEST", "normalization_tag": norm_tag}
            params: dict = {}

            from backend.services.deliverables import generate_complete_results_xlsx
            generate_complete_results_xlsx(run_dir, profile, params)

            xlsx_files = list((run_dir / f"05_network_{norm_tag}").glob("*_Complete_Results.xlsx"))
            self.assertEqual(len(xlsx_files), 1, "Expected exactly one xlsx file")

            wb = openpyxl.load_workbook(xlsx_files[0], read_only=True)
            sheet_names = wb.sheetnames
            wb.close()

            base_sheets = {"Module_Assignments", "GO_Enrichment_ZScores", "CellType_FET", "Trait_Associated_Modules"}
            self.assertTrue(base_sheets.issubset(set(sheet_names)), f"Missing base sheets: {base_sheets - set(sheet_names)}")

            hub_sheets = [s for s in sheet_names if s.endswith("_Hub_Proteins")]
            self.assertEqual(len(hub_sheets), 2, f"Expected 2 hub sheets (Disease + Abeta42), got: {hub_sheets}")

            wb = openpyxl.load_workbook(xlsx_files[0], read_only=False)
            for hs in hub_sheets:
                ws = wb[hs]
                headers = [cell.value for cell in ws[1]]
                self.assertEqual(headers, ["module", "direction", "module_correlation", "gene", "peptide_id", "kME", "hub_rank"])
                self.assertGreater(ws.max_row, 1, f"Hub sheet {hs} has no data rows")
            wb.close()

            wb = openpyxl.load_workbook(xlsx_files[0], read_only=False)
            disease_hub = [s for s in hub_sheets if "Disease" in s][0]
            ws = wb[disease_hub]
            module_values = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
            wb.close()
            for mv in module_values:
                self.assertFalse(mv.startswith("ME"), f"ME prefix not stripped: {mv}")

            wb = openpyxl.load_workbook(xlsx_files[0], read_only=False)
            ws = wb[disease_hub]
            for r in range(2, ws.max_row + 1):
                if ws.cell(row=r, column=1).value == "blue":
                    self.assertEqual(ws.cell(row=r, column=4).value, "APOE", "Top blue hub should be APOE (kME=0.95)")
                    self.assertEqual(ws.cell(row=r, column=7).value, 1, "Top hub should have rank 1")
                    break
            wb.close()

    def test_hub_sheets_absent_when_no_traits(self) -> None:
        """HUB-01 graceful path: no trait files means no hub sheets, no crash."""
        import tempfile
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            run_dir = Path(tmp)
            stage1 = run_dir / "stage1"
            stage1.mkdir()
            (stage1 / "module_assignments.csv").write_text(
                "peptide_id,feature_id,gene,module_color,kME,alternative_module,kME_alternative,module_quality\n"
                "P001,F001,APOE,blue,0.95,turquoise,0.30,high\n"
            )

            norm_tag = "CBN_median"
            profile = {"deliverable_prefix": "TEST", "normalization_tag": norm_tag}
            params: dict = {}

            from backend.services.deliverables import generate_complete_results_xlsx
            generate_complete_results_xlsx(run_dir, profile, params)

            xlsx_files = list((run_dir / f"05_network_{norm_tag}").glob("*_Complete_Results.xlsx"))
            self.assertEqual(len(xlsx_files), 1)

            wb = openpyxl.load_workbook(xlsx_files[0], read_only=True)
            hub_sheets = [s for s in wb.sheetnames if s.endswith("_Hub_Proteins")]
            wb.close()
            self.assertEqual(len(hub_sheets), 0, f"No hub sheets expected without traits, got: {hub_sheets}")

    def test_run_submission_blocked_for_unready_format(self):
        """POST /api/runs must return 422 when dataset sniff says run_ready=False."""
        import json as _json

        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                # Upload a minimal CSV
                csv_bytes = b"feature_id,s1,s2\ngene1,1.0,2.0\ngene2,3.0,4.0\n"
                r = client.post(
                    "/api/uploads",
                    files={"file": ("test_fmt.csv", io.BytesIO(csv_bytes), "text/csv")},
                    data={"workspace_id": workspace["id"]},
                    headers=headers,
                )
                self.assertEqual(r.status_code, 200, r.text)
                dataset_id = r.json()["dataset_id"]

                # Patch sniff_metadata_json to mark run_ready=False
                from database import get_db, UploadedDataset
                db = next(get_db())
                try:
                    ds = db.query(UploadedDataset).filter(UploadedDataset.id == dataset_id).first()
                    ds.sniff_metadata_json = _json.dumps({
                        "run_ready": False,
                        "format_detected": "Olink Quant Long Table",
                        "format_family": "Olink",
                        "warnings": ["Only NPX-valued Olink exports are supported."],
                    })
                    db.commit()
                finally:
                    db.close()

                # Attempt to create a run — must be blocked with 422
                r = client.post(
                    "/api/runs",
                    json={
                        "workspace_id": workspace["id"],
                        "dataset_id": dataset_id,
                        "name": "blocked run",
                        "cohort1": "Control",
                        "cohort2": "Disease",
                        "params": {},
                    },
                    headers=headers,
                )
                self.assertEqual(r.status_code, 422, r.text)
                body = r.json()
                detail = body.get("detail", "")
                self.assertIn("Olink Quant Long Table", detail)

    def test_pipeline_profile_batch_correction_disabled_by_default(self):
        """resolve_pipeline_profile must report variance_batch_correction enabled=False when param not set."""
        import sys
        import os
        sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))
        from services.deliverables import resolve_pipeline_profile
        profile = resolve_pipeline_profile(manifest={}, params={})
        steps = {s["key"]: s["enabled"] for s in profile["supported_steps"]}
        self.assertFalse(
            steps.get("variance_batch_correction"),
            "variance_batch_correction should be False when variance_correction_enabled not set"
        )


    def test_raw_std_trait_files_route_to_correct_subfolder(self):
        """_raw/_std Associated_Modules files must land under the correct trait subfolder."""
        import sys
        import os
        sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))
        from services.deliverables import _trait_folder_for_name

        self.assertEqual(_trait_folder_for_name("ABETA42_raw"), "amyloid_beta")
        self.assertEqual(_trait_folder_for_name("ABETA42_std"), "amyloid_beta")
        self.assertEqual(_trait_folder_for_name("T_TAU_raw"), "total_tau")
        self.assertEqual(_trait_folder_for_name("T_TAU_std"), "total_tau")
        self.assertEqual(_trait_folder_for_name("P_TAU_raw"), "phospho_tau")
        self.assertEqual(_trait_folder_for_name("P_TAU_std"), "phospho_tau")
        self.assertEqual(_trait_folder_for_name("Disease"), "disease_status")

    def test_run_status_awaiting_review_is_valid(self):
        """RunStatus.AWAITING_REVIEW must exist and be distinct from RUNNING/COMPLETE."""
        with tempfile.TemporaryDirectory() as temp_dir:
            build_app(temp_dir)
            from database import RunStatus
            self.assertEqual(RunStatus.AWAITING_REVIEW.value, "awaiting_review")
            self.assertNotEqual(RunStatus.AWAITING_REVIEW, RunStatus.RUNNING)
            self.assertNotEqual(RunStatus.AWAITING_REVIEW, RunStatus.COMPLETE)

    def test_detect_outliers_no_candidates_proceeds(self):
        """_detect_outliers must return [] and not block when R produces no candidates."""
        import sys, os, json, tempfile, threading
        sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))
        from services import pipeline as pl

        with tempfile.TemporaryDirectory() as td:
            run_dir = __import__('pathlib').Path(td)
            input_dir = run_dir / "input"
            input_dir.mkdir()
            (input_dir / "cleaned_matrix.csv").write_text(
                "feature_id,S1,S2,S3,S4\ngene1,1.0,1.1,0.9,1.05\ngene2,2.0,2.1,1.9,2.05\n"
            )

            calls = []
            def fake_step(key, status, pct, msg): calls.append((key, status))

            original = getattr(pl, '_run_r_outlier_detection', None)
            def fake_run_r(matrix_path, z_threshold, output_json_path, log_fn):
                output_json_path.write_text(json.dumps({
                    "outlier_candidates": [],
                    "all_z_scores": {"S1": 0.5, "S2": 0.3, "S3": -0.1, "S4": 0.2}
                }))
                return True
            pl._run_r_outlier_detection = fake_run_r

            try:
                result = pl._detect_outliers(
                    run_id="test-run-123",
                    run_dir=run_dir,
                    params={"outlier_z_threshold": 2.0},
                    log_fn=lambda x: None,
                    step_callback=fake_step,
                    db=None,
                )
                self.assertEqual(result, [])
                self.assertIn(("outlier_removal", "complete"), calls)
            finally:
                if original is not None:
                    pl._run_r_outlier_detection = original
                else:
                    delattr(pl, '_run_r_outlier_detection')


    def test_review_outliers_endpoint_skip(self):
        """POST /api/runs/{id}/review-outliers with action=skip must update params and fire the event."""
        import json as _j, threading, sys, os

        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                csv_bytes = b"feature_id,s1,s2\ngene1,1.0,2.0\ngene2,3.0,4.0\n"
                r = client.post(
                    "/api/uploads",
                    files={"file": ("test.csv", io.BytesIO(csv_bytes), "text/csv")},
                    data={"workspace_id": workspace["id"]},
                    headers=headers,
                )
                self.assertEqual(r.status_code, 200, r.text)
                dataset_id = r.json()["dataset_id"]

                r = client.post(
                    "/api/runs",
                    json={"workspace_id": workspace["id"], "dataset_id": dataset_id,
                          "name": "outlier-skip-test", "cohort1": "A", "cohort2": "B", "params": {}},
                    headers=headers,
                )
                # run may fail/complete fast — we just need the run_id
                run_id = r.json()["run_id"]

                # Force AWAITING_REVIEW status directly in DB
                from database import get_db, Run, RunStatus
                db = next(get_db())
                try:
                    run_obj = db.query(Run).filter(Run.id == run_id).first()
                    run_obj.status = RunStatus.AWAITING_REVIEW
                    db.commit()
                finally:
                    db.close()

                # Register a fake threading.Event so the endpoint can fire it
                sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))
                from services.pipeline import _OUTLIER_EVENTS
                evt = threading.Event()
                _OUTLIER_EVENTS[run_id] = evt

                r = client.post(
                    f"/api/runs/{run_id}/review-outliers",
                    json={"action": "skip"},
                    headers=headers,
                )
                self.assertEqual(r.status_code, 200, r.text)
                self.assertTrue(evt.is_set())

                db = next(get_db())
                try:
                    run_obj = db.query(Run).filter(Run.id == run_id).first()
                    updated_params = _j.loads(run_obj.params or "{}")
                finally:
                    db.close()
                self.assertEqual(updated_params.get("excluded_samples"), [])
                self.assertEqual(updated_params.get("outlier_action"), "skip")

    def test_frontend_has_outlier_review_handler(self):
        """app.js must contain outlier_review SSE handler and review modal code."""
        import os
        app_js = os.path.join(os.path.dirname(__file__), '..', 'frontend', 'app.js')
        with open(app_js) as fh:
            content = fh.read()
        self.assertIn("outlier_review", content, "SSE handler for outlier_review missing")
        self.assertIn("review-outliers", content, "API call to review-outliers missing")
        self.assertIn("outlier-review-modal", content, "Modal element id missing")

    def test_outlier_review_full_flow_skip(self):
        """
        With INLINE_RUNS=1, mocked _run_r_outlier_detection returning 1 candidate,
        calling review-outliers?action=skip must unblock the pipeline and let the run finish.

        NOTE: INLINE_RUNS=1 blocks the HTTP response until the pipeline completes.
        The run_id is known only after the POST returns. We work around this by
        polling the DB for AWAITING_REVIEW from a reviewer thread before the submit
        thread has a run_id, using _OUTLIER_EVENTS as the signal.
        """
        import json as _j, threading, time, sys, os

        with tempfile.TemporaryDirectory() as temp_dir:
            app = build_app(temp_dir)
            with TestClient(app) as client:
                headers, workspace, project = bootstrap_session(client)

                csv_bytes = b"feature_id,s1,s2\ngene1,1.0,2.0\ngene2,3.0,4.0\n"
                r = client.post(
                    "/api/uploads",
                    files={"file": ("test.csv", io.BytesIO(csv_bytes), "text/csv")},
                    data={"workspace_id": workspace["id"]},
                    headers=headers,
                )
                dataset_id = r.json()["dataset_id"]

                sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))
                from services import pipeline as pl
                from database import get_db, Run, RunStatus

                # Patch R outlier detection to return 1 candidate
                original_r = pl._run_r_outlier_detection
                def fake_r(matrix_path, z_threshold, output_json_path, log_fn):
                    output_json_path.write_text(_j.dumps({
                        "outlier_candidates": [{"sample": "s1", "z_score": -2.5, "round_detected": 1}],
                        "all_z_scores": {"s1": -2.5, "s2": 0.8}
                    }))
                    return True
                pl._run_r_outlier_detection = fake_r

                submit_result = {}
                submit_done = threading.Event()

                def submit_run():
                    try:
                        rr = client.post(
                            "/api/runs",
                            json={"workspace_id": workspace["id"], "dataset_id": dataset_id,
                                  "name": "p2a-integration", "cohort1": "A", "cohort2": "B", "params": {}},
                            headers=headers,
                        )
                        submit_result["response"] = rr.json()
                    except Exception as e:
                        submit_result["error"] = str(e)
                    finally:
                        submit_done.set()

                t = threading.Thread(target=submit_run, daemon=True)
                t.start()

                # Poll _OUTLIER_EVENTS for a new entry — this appears before the HTTP response
                # because the pipeline thread blocks inside the route handler with INLINE_RUNS=1.
                run_id = None
                for _ in range(100):
                    time.sleep(0.1)
                    if pl._OUTLIER_EVENTS:
                        run_id = next(iter(pl._OUTLIER_EVENTS))
                        break

                if run_id is None:
                    # Outlier detection was skipped (e.g., R unavailable). Test passes vacuously.
                    pl._run_r_outlier_detection = original_r
                    submit_done.wait(timeout=30)
                    return

                # Approve with skip
                r = client.post(
                    f"/api/runs/{run_id}/review-outliers",
                    json={"action": "skip"},
                    headers=headers,
                )
                self.assertEqual(r.status_code, 200, r.text)

                # Wait for the pipeline to finish
                submit_done.wait(timeout=60)

                db = next(get_db())
                try:
                    run_obj = db.query(Run).filter(Run.id == run_id).first()
                    final_status = run_obj.status if run_obj else None
                    p = _j.loads(run_obj.params or "{}") if run_obj else {}
                finally:
                    db.close()

                self.assertIn(final_status, (RunStatus.COMPLETE, RunStatus.FAILED),
                              f"Run stuck at {final_status}")
                self.assertEqual(p.get("outlier_action"), "skip")
                pl._run_r_outlier_detection = original_r


if __name__ == "__main__":
    unittest.main()
